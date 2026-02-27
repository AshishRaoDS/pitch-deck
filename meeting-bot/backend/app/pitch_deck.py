"""
Pitch deck generation using GPT-4o for content and python-pptx for rendering.

Flow:
1. Send transcript + enriched topics to GPT-4o → structured slide content.
2. Use python-pptx to build a .pptx file from the structured content.
3. Return the file path.
"""

import io
import json
import os
import textwrap
import uuid

from openai import AsyncOpenAI
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from .config import settings

# ---------------------------------------------------------------------------
# Constants / theme
# ---------------------------------------------------------------------------

THEME = {
    "bg": RGBColor(0x0F, 0x17, 0x2A),       # deep navy
    "accent": RGBColor(0x4F, 0x8E, 0xFF),    # bright blue
    "title_text": RGBColor(0xFF, 0xFF, 0xFF),
    "body_text": RGBColor(0xD0, 0xD8, 0xEA),
    "highlight": RGBColor(0x4F, 0xE3, 0xC0), # teal
}

OUTPUT_DIR = "/tmp/meeting-bot-decks"
os.makedirs(OUTPUT_DIR, exist_ok=True)

# ---------------------------------------------------------------------------
# GPT-4o: generate slide content
# ---------------------------------------------------------------------------

_SYSTEM_PROMPT = """\
You are a professional pitch deck writer. Given a meeting transcript and a list
of enriched topics (each with web-search results), produce structured content
for a compelling pitch deck.

Return a JSON object with this exact shape:
{
  "deck_title": "<catchy overall title for the pitch>",
  "slides": [
    {
      "type": "title",
      "title": "<deck title>",
      "subtitle": "<one-line value proposition>"
    },
    {
      "type": "agenda",
      "title": "Agenda",
      "bullets": ["<item 1>", "<item 2>", ...]
    },
    {
      "type": "topic",
      "title": "<topic name>",
      "bullets": ["<key insight 1>", "<key insight 2>", ...],
      "source": "<one of the search result URLs, if relevant>"
    },
    ...more topic slides...,
    {
      "type": "summary",
      "title": "Key Takeaways",
      "bullets": ["<takeaway 1>", "<takeaway 2>", ...]
    },
    {
      "type": "cta",
      "title": "Next Steps",
      "bullets": ["<action 1>", "<action 2>", ...]
    }
  ]
}

Rules:
- Include one "topic" slide per enriched topic (3-8 slides total for topics).
- Each bullet should be 10-20 words, punchy and specific.
- Only return valid JSON – no markdown fences, no extra text.
"""


async def _generate_slide_content(transcript: str, enriched_topics: list[dict]) -> dict:
    client = AsyncOpenAI(api_key=settings.openai_api_key)

    topics_text = json.dumps(enriched_topics, indent=2)
    user_msg = (
        f"Meeting transcript:\n{transcript}\n\n"
        f"Enriched topics (with web search results):\n{topics_text}"
    )

    response = await client.chat.completions.create(
        model="gpt-4o",
        response_format={"type": "json_object"},
        messages=[
            {"role": "system", "content": _SYSTEM_PROMPT},
            {"role": "user", "content": user_msg},
        ],
        temperature=0.4,
    )

    content = response.choices[0].message.content or "{}"
    return json.loads(content)


# ---------------------------------------------------------------------------
# python-pptx rendering helpers
# ---------------------------------------------------------------------------

def _set_slide_background(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _add_text_box(slide, text: str, left, top, width, height,
                  font_size: int, bold: bool, color: RGBColor,
                  align=PP_ALIGN.LEFT, wrap: bool = True) -> None:
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color


def _render_title_slide(prs: Presentation, slide_data: dict) -> None:
    slide_layout = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(slide_layout)
    _set_slide_background(slide, THEME["bg"])

    W, H = prs.slide_width, prs.slide_height
    pad = Inches(0.8)

    # Accent bar at top
    bar = slide.shapes.add_shape(1, 0, 0, W, Inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = THEME["accent"]
    bar.line.fill.background()

    _add_text_box(slide, slide_data.get("title", "Meeting Pitch"),
                  pad, H * 0.3, W - 2 * pad, Inches(1.5),
                  font_size=40, bold=True, color=THEME["title_text"],
                  align=PP_ALIGN.CENTER)

    _add_text_box(slide, slide_data.get("subtitle", ""),
                  pad, H * 0.55, W - 2 * pad, Inches(1),
                  font_size=20, bold=False, color=THEME["highlight"],
                  align=PP_ALIGN.CENTER)


def _render_bullet_slide(prs: Presentation, slide_data: dict, accent_color: RGBColor | None = None) -> None:
    accent_color = accent_color or THEME["accent"]
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    _set_slide_background(slide, THEME["bg"])

    W, H = prs.slide_width, prs.slide_height
    pad = Inches(0.6)

    # Title
    _add_text_box(slide, slide_data.get("title", ""),
                  pad, Inches(0.35), W - 2 * pad, Inches(0.8),
                  font_size=28, bold=True, color=THEME["title_text"])

    # Divider line
    line = slide.shapes.add_shape(1, pad, Inches(1.25), W - 2 * pad, Inches(0.04))
    line.fill.solid()
    line.fill.fore_color.rgb = accent_color
    line.line.fill.background()

    # Bullets
    bullets: list[str] = slide_data.get("bullets", [])
    y = Inches(1.4)
    bullet_h = Inches(0.55)
    for bullet in bullets[:7]:
        dot = slide.shapes.add_shape(1, pad, y + Inches(0.14), Inches(0.12), Inches(0.12))
        dot.fill.solid()
        dot.fill.fore_color.rgb = accent_color
        dot.line.fill.background()

        _add_text_box(slide, bullet,
                      pad + Inches(0.28), y, W - 2 * pad - Inches(0.28), bullet_h,
                      font_size=16, bold=False, color=THEME["body_text"])
        y += bullet_h

    # Optional source URL (small, bottom right)
    source = slide_data.get("source", "")
    if source:
        _add_text_box(slide, f"Source: {source}",
                      pad, H - Inches(0.45), W - 2 * pad, Inches(0.35),
                      font_size=9, bold=False, color=RGBColor(0x80, 0x90, 0xAA),
                      align=PP_ALIGN.RIGHT)


def _render_cta_slide(prs: Presentation, slide_data: dict) -> None:
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    _set_slide_background(slide, THEME["accent"])

    W, H = prs.slide_width, prs.slide_height
    pad = Inches(0.8)

    _add_text_box(slide, slide_data.get("title", "Next Steps"),
                  pad, Inches(0.5), W - 2 * pad, Inches(0.9),
                  font_size=32, bold=True, color=THEME["bg"])

    bullets: list[str] = slide_data.get("bullets", [])
    y = Inches(1.5)
    for bullet in bullets[:6]:
        _add_text_box(slide, f"→  {bullet}",
                      pad, y, W - 2 * pad, Inches(0.6),
                      font_size=17, bold=False, color=THEME["bg"])
        y += Inches(0.65)


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

async def generate_pitch_deck(transcript: str, enriched_topics: list[dict]) -> str:
    """
    Generate a .pptx pitch deck.

    Parameters
    ----------
    transcript       : Full meeting transcript text.
    enriched_topics  : Output of search.enrich_topics().

    Returns
    -------
    Absolute path to the generated .pptx file.
    """
    deck_data = await _generate_slide_content(transcript, enriched_topics)

    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    for slide in deck_data.get("slides", []):
        slide_type = slide.get("type", "bullet")

        if slide_type == "title":
            _render_title_slide(prs, slide)
        elif slide_type == "cta":
            _render_cta_slide(prs, slide)
        else:
            accent = THEME["highlight"] if slide_type == "summary" else THEME["accent"]
            _render_bullet_slide(prs, slide, accent_color=accent)

    filename = f"pitch_{uuid.uuid4().hex[:8]}.pptx"
    filepath = os.path.join(OUTPUT_DIR, filename)
    prs.save(filepath)
    return filepath
